"""Doi chieu build voi plan tuan cua QA: build khac yeu cau trong plan -> preflight LOI, dung luot.

User chot 30/10/2026: AIP882 4.3.0 plan ghi "- SDK 3.6.1" nhung build in SDK 3.7.0, tool chay checklist
Next-Gen vo nghia. Tu nay khac plan thi dung lai, bao QA o qa-auto (bot dang ket qua DUNG duoi tin build).

Plan la Google Sheet (link trong bien `RC_PLAN_TUAN_SHEET`, khong ghi vao repo public), moi tuan mot tab ten
`<ddmm>-<ddmmyyyy>` (`2610-30102026` = 26/10 - 30/10/2026; tab cu thang 1 chu so nhu `108 - 1482026`
khong doc - chi can tab tuan hien tai). Dong cua app:
o ma app (`AIP882`, `AIP882 (từ 3.6.0)`) + o noi dung (`- SDK 3.6.1`, `- Apply SDK Next Gen`). Dong QA
(`test v4.3.0 doing`) khong co so SDK nen tu bo qua.
"""

from __future__ import annotations

import datetime as dt
import re

from . import build_ma_app, sdk_probe, tc_catalog, tc_source
from .models import RcError

ENV = "RC_PLAN_TUAN_SHEET"
TAB_RE = re.compile(r"^\s*(\d{1,2})(\d{2})\s*-\s*(\d{1,2})(\d{2})(\d{4})\s*$")


def tuan_cua_tab(ten: str) -> tuple[dt.date, dt.date] | None:
    """'2610-30102026' -> (26/10/2026, 30/10/2026); ten khong dung mau -> None."""
    m = TAB_RE.match(ten or "")
    if not m:
        return None
    d1, m1, d2, m2, nam = (int(x) for x in m.groups())
    try:
        cuoi = dt.date(nam, m2, d2)
        return dt.date(nam - 1 if m1 > m2 else nam, m1, d1), cuoi   # tuan vat qua nam moi
    except ValueError:
        return None


def ma_cua_package(path, package: str) -> str:
    """Ten tab (ma app) co dong "Package name" = `package` trong sheet Check thong so KT; "" khi khong co."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            for r in ws.iter_rows(values_only=True):
                o = [str(v).strip() for v in r if v not in (None, "")]
                if len(o) >= 2 and o[0].lower() == "package name" and o[1] == package:
                    return ws.title.strip().upper()
        return ""
    finally:
        wb.close()


def yeu_cau(path, ma: str, ngay: dt.date) -> dict | None:
    """Yeu cau SDK cua app `ma` trong tab tuan chua `ngay`: {tab, sdk: [X.Y.Z], next_gen, dong: [noi dung]}.
    Khong co tab tuan -> None; co tab nhung app khong co dong ghi SDK -> sdk rong, next_gen False."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = next((w for w in wb.worksheets if (t := tuan_cua_tab(w.title)) and t[0] <= ngay <= t[1]), None)
        if ws is None:
            return None
        ra = {"tab": ws.title.strip(), "sdk": [], "next_gen": False, "dong": []}
        for r in ws.iter_rows(values_only=True):
            o = [str(v).strip() for v in r if isinstance(v, str) and v.strip()]
            if not any((m := build_ma_app.MA_RE.match(c)) and m.group(1).upper() == ma for c in o):
                continue
            for c in o:
                so = [s for s in build_ma_app.SDK_RE.findall(c) if s not in ra["sdk"]]
                ng = bool(build_ma_app.NEXT_GEN_RE.search(c)) and not build_ma_app.REVERT_RE.search(c)
                if so or ng:
                    ra["sdk"] += so
                    ra["next_gen"] = ra["next_gen"] or ng
                    ra["dong"].append(" ".join(c.split()))
        return ra
    finally:
        wb.close()


def lech(yc: dict, sdk_khai: str, sdk_log: str) -> list[str]:
    """Cac diem build khac plan (rong = khop). So theo X.Y.Z, bo hau to -alphaNN."""
    ra = []
    can = {tc_catalog.parse_version(s) for s in yc["sdk"]}
    for nguon, sdk in (("tin build ghi", sdk_khai), ("app in log", sdk_log)):
        v = tc_catalog.parse_version(sdk)
        if v is None:
            continue
        if can and v not in can:
            ra.append(f"{nguon} SDK {sdk}, plan ghi {' / '.join(yc['sdk'])}")
        elif not can and yc["next_gen"] and not sdk_probe.la_nextgen(sdk):
            ra.append(f"{nguon} SDK {sdk}, plan ghi Next-Gen")
    return ra


def tai(link: str, out_dir) -> str:
    try:
        return tc_source.tai_ve(link, out_dir)
    except RcError as exc:
        raise RcError(f"không tải được plan tuần: {exc}") from None

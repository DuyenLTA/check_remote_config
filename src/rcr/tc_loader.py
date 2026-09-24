"""Doc file testcase XLSX -> list[Case]. Thuan du lieu, khong can device.

MAP COT THEO TEN, KHONG THEO CHI SO. File TC cua du an khac co the them/bo cot
(vd them "Priority", "Note") -> hardcode index la parse lech het. Cot la thi bao
loi RO, khong doan.

Dong nhom (chi co o dau, khong phai so) bi bo: file that dung dong do lam tieu
de nhom, vd "Tab Moment (301) - Config / Feature Flag".

O TRONG O `Feature`/`Test Description` KE THUA dong tren: file that de trong khi
lap lai, chi ghi o dong dau cua nhom.

CHON SHEET THEO HEADER, KHONG THEO TEN: bo TC dung chung cho moi app (moi ban
SDK mot bo) dat ten sheet tuy y. Uu tien sheet 'Test Cases' neu co, khong thi
lay sheet dau tien tim thay header.
"""

from __future__ import annotations

import re
from pathlib import Path

from .models import Case, RcError

SHEET_NAME = "Test Cases"

# Ten cot can, sau khi chuan hoa (lower + gop khoang trang). Moi muc la cac
# bien the chap nhan duoc cua cung mot cot.
COLUMNS = {
    "n": ("n°", "no", "no.", "stt", "#", "id"),
    "feature": ("feature", "tinh nang"),
    "description": ("test description", "description", "mo ta"),
    "sub_scenario": ("sub-scenario", "sub scenario", "scenario", "kich ban"),
    "precondition": ("precondition", "pre-condition", "dieu kien truoc"),
    "test_data": ("test data", "data", "du lieu"),
    "actions": ("action", "actions", "steps", "thao tac"),
    "expects": ("expected result", "expected", "ket qua mong doi"),
}
REQUIRED = ("n", "actions", "expects")
# Key RC nam o Test Data HOAC Precondition (bo TC tu SDK 3.2.0 bo han cot Test Data,
# ghi key trong Precondition) -> can it nhat mot trong hai.
KEY_COLUMNS = ("test_data", "precondition")

# Buoc duoc danh so: "1. ...", "2) ...". Tach theo dau dong.
STEP_SPLIT = re.compile(r"(?m)^\s*\d+\s*[.)]\s*")
CASE_NO = re.compile(r"(\d+)(?:\.0+)?")
HEADER_SCAN_ROWS = 12  # header cua file that nam o dong 2; quet du rong


def load(path: str | Path) -> list[Case]:
    """Doc sheet chua bang testcase. Raise RcError kem thong bao doc duoc."""
    import openpyxl

    p = Path(path)
    if not p.is_file():
        raise RcError(f"Khong thay file: {p}")
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl nem nhieu loai
        raise RcError(f"Khong doc duoc {p.name}: {exc}") from exc
    try:
        names = sorted(wb.sheetnames, key=lambda n: n != SHEET_NAME)
        sheets = [(n, [tuple(r) for r in wb[n].iter_rows(values_only=True)]) for n in names]
    finally:
        wb.close()
    first_err = None
    for name, rows in sheets:
        try:
            _find_header(rows, name)
        except RcError as exc:
            first_err = first_err or exc  # giu chi tiet cot thieu cua sheet uu tien
            continue
        return parse_rows(rows, f"{p.name} / {name}")
    raise RcError(
        f"{p.name}: khong sheet nao co dong header testcase "
        f"(sheet co trong file: {', '.join(wb.sheetnames)}).\n{first_err}"
    )


def parse_rows(rows: list[tuple], source: str = "") -> list[Case]:
    """Tach khoi `load` de test bang bang du lieu, khong can file that."""
    idx, header_at = _find_header(rows, source)
    out: list[Case] = []
    inherit = {"feature": "", "description": ""}
    for row in rows[header_at + 1 :]:
        cells = [_text(c) for c in row]
        n = _case_no(_cell(cells, idx, "n"))
        if not n:
            continue  # dong nhom / dong trong / dong tong
        for key in ("feature", "description"):
            val = _cell(cells, idx, key)
            if val:
                inherit[key] = val
        out.append(
            Case(
                n=n,
                feature=inherit["feature"],
                description=inherit["description"],
                sub_scenario=_cell(cells, idx, "sub_scenario"),
                precondition=_cell(cells, idx, "precondition"),
                test_data=_cell(cells, idx, "test_data"),
                actions=split_steps(_cell(cells, idx, "actions")),
                expects=split_steps(_cell(cells, idx, "expects")),
            )
        )
    if not out:
        raise RcError(
            f"{source or 'File'}: tim thay header nhung khong co dong case nao "
            f"(cot {_label(idx, 'n')!r} phai la so)."
        )
    return out


def split_steps(text: str) -> tuple[str, ...]:
    """'1. mo tab\\n2. quan sat' -> ('mo tab', 'quan sat').

    Khong danh so -> tra ve mot buoc duy nhat la ca doan.
    """
    if not text.strip():
        return ()
    parts = [s.strip() for s in STEP_SPLIT.split(text) if s.strip()]
    return tuple(parts) if parts else (text.strip(),)


def _find_header(rows: list[tuple], source: str) -> tuple[dict[str, int], int]:
    """Quet may dong dau tim dong header -> ({ten_cot: chi_so}, so_dong)."""
    best: tuple[dict[str, int], int] | None = None
    for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        idx = _match_header(row)
        if all(k in idx for k in REQUIRED) and any(k in idx for k in KEY_COLUMNS):
            return idx, i
        if idx and (best is None or len(idx) > len(best[0])):
            best = (idx, i)
    found = ", ".join(sorted(best[0])) if best else "(khong nhan ra cot nao)"
    raise RcError(
        f"{source or 'File'}: khong tim thay dong header trong {HEADER_SCAN_ROWS} dong dau.\n"
        f"Can du cac cot: {', '.join(REQUIRED)} + mot trong {' / '.join(KEY_COLUMNS)}. "
        f"Nhan ra duoc: {found}.\n"
        "Kiem tra ten cot trong sheet testcase."
    )


def _match_header(row: tuple) -> dict[str, int]:
    idx: dict[str, int] = {}
    for i, cell in enumerate(row):
        norm = _norm(_text(cell))
        if not norm:
            continue
        for key, aliases in COLUMNS.items():
            if key not in idx and norm in aliases:
                idx[key] = i
                break
    return idx


def _case_no(text: str) -> str:
    """'7' / '7.0' -> '7'. Google Sheet export so thanh float nen N° ra '7.0'.

    Khong phai so nguyen (dong nhom, 'TC-01', '1.5') -> ''.
    """
    m = CASE_NO.fullmatch(text)
    return m.group(1) if m else ""


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip().lower()


def _text(cell) -> str:
    return "" if cell is None else str(cell).strip()


def _cell(cells: list[str], idx: dict[str, int], key: str) -> str:
    i = idx.get(key)
    return cells[i] if i is not None and i < len(cells) else ""


def _label(idx: dict[str, int], key: str) -> str:
    return key if key in idx else "?"

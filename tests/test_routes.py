"""Test route qua TestClient. Khong can device: cam FakeAdb vao app_state.

Diem phai gac:
  - chi bind 127.0.0.1: Host khac -> 403
  - loi cua tang duoi ra 400 kem thong bao doc duoc, KHONG phai 500 + traceback
  - upload file testcase -> loc bang whitelist = key THAT cua app
  - file sai duoi / file rong / chua doc baseline -> 400
"""

from __future__ import annotations

import asyncio

import pytest
from conftest import PKG, SERIAL
from fastapi.testclient import TestClient

from rcr import app_state, rc_baseline
from rcr.main import app


@pytest.fixture
def client(adb, monkeypatch):
    monkeypatch.setattr(app_state, "client", lambda: adb)
    app_state._baselines.clear()
    # TestClient mac dinh gui `Host: testserver` -> middleware chan dung (403).
    # Dat base_url ve 127.0.0.1 de test duoc cac route.
    yield TestClient(app, base_url="http://127.0.0.1")
    app_state._baselines.clear()


@pytest.fixture
def with_baseline(client, adb):
    app_state.put_baseline(asyncio.run(rc_baseline.read(adb, SERIAL, PKG)))
    return client


def make_xlsx(path, rows):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"
    ws.append(("N°", "Feature", "Test Description", "Sub-scenario", "Precondition",
               "Test Data", "Action", "Expected Result", "PASS/FAIL"))
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def upload(client, path, serial=SERIAL, package=PKG):
    with open(path, "rb") as fh:
        return client.post(
            f"/api/testcases?serial={serial}&package={package}",
            files={"file": (path.name, fh,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )


# ---- bao ve --------------------------------------------------------------


def test_host_khac_bi_chan_403(client):
    r = client.get("/api/devices", headers={"Host": "evil.local"})
    assert r.status_code == 403


def test_localhost_duoc_phep(client):
    assert client.get("/api/devices").status_code == 200


def test_trang_chu_tra_html(client):
    r = client.get("/")
    assert r.status_code == 200
    assert "text/html" in r.headers["content-type"]


# ---- baseline ------------------------------------------------------------


def test_doc_baseline(client):
    r = client.post(f"/api/baseline?serial={SERIAL}&package={PKG}")
    assert r.status_code == 200
    s = r.json()
    assert s["keys"] == 6
    assert s["mirrored_keys"] == 4
    assert s["write_mode"] == "run-as"


def test_app_khong_debuggable_ra_400_khong_phai_500(client, adb):
    adb.debuggable = False
    adb.rooted = False
    r = client.post(f"/api/baseline?serial={SERIAL}&package={PKG}")
    assert r.status_code == 400
    assert "not debuggable" in r.json()["detail"]


def test_chua_doc_baseline_thi_400(client):
    r = client.get(f"/api/baseline/keys?serial={SERIAL}&package={PKG}")
    assert r.status_code == 400
    assert "Chua doc baseline" in r.json()["detail"]


def test_tim_key(with_baseline):
    r = with_baseline.get(f"/api/baseline/keys?serial={SERIAL}&package={PKG}&q=onb3")
    assert r.status_code == 200
    body = r.json()
    assert body["total"] == 1
    assert body["keys"][0]["key"] == "enable_onb3_screen"
    assert body["keys"][0]["mirrored"] is True


# ---- upload file testcase -------------------------------------------------


def test_nap_file_testcase_loc_bang_whitelist(with_baseline, tmp_path):
    """`click_area` khong thuoc RC cua app -> phai bi loai, khong can khai gi."""
    p = make_xlsx(tmp_path / "tc.xlsx", [
        ("Nhom A - Config",) + ("",) * 8,
        ("1", "Tab X", "Config", "tat onb3", "", "enable_onb3_screen = false",
         "1. Mo app.", "1. Man onb3 khong hien.", ""),
        ("2", "", "", "param analytics", "", "click_area = cta_button",
         "1. Nhan banner.", "1. Event ban dung.", ""),
        ("3", "", "", "runtime toggle", "", "enable_onb3_screen: true → false → true",
         "1. Doi khi dang chay.", "1. Doi ngay.", ""),
    ])
    r = upload(with_baseline, p)
    assert r.status_code == 200, r.text
    b = r.json()
    assert b["total"] == 3
    assert b["runnable"] == 1
    assert b["whitelist_size"] == 6

    by_n = {c["n"]: c for c in b["cases"]}
    assert by_n["1"]["overrides"] == {"enable_onb3_screen": "false"}
    assert by_n["1"]["runnable"] is True
    assert by_n["1"]["steps"] == 1 and by_n["1"]["assertions"] == 1

    assert by_n["2"]["overrides"] == {}
    assert by_n["2"]["ignored"] == ["click_area"]
    assert "khong co key nao thuoc remote config" in by_n["2"]["needs_human"]

    assert by_n["3"]["overrides"] == {}
    assert "runtime toggle" in by_n["3"]["needs_human"]


def test_nap_file_thieu_cot_thi_400(with_baseline, tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"
    ws.append(("N°", "Feature"))
    ws.append(("1", "x"))
    p = tmp_path / "thieu.xlsx"
    wb.save(p)
    r = upload(with_baseline, p)
    assert r.status_code == 400
    assert "khong tim thay dong header" in r.json()["detail"]


def test_nap_file_thieu_sheet_thi_400(with_baseline, tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "Sheet khac"
    p = tmp_path / "sai-sheet.xlsx"
    wb.save(p)
    r = upload(with_baseline, p)
    assert r.status_code == 400
    assert "Test Cases" in r.json()["detail"]


def test_file_sai_duoi_thi_400(with_baseline, tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("khong phai xlsx")
    with open(p, "rb") as fh:
        r = with_baseline.post(
            f"/api/testcases?serial={SERIAL}&package={PKG}",
            files={"file": ("a.txt", fh, "text/plain")},
        )
    assert r.status_code == 400
    assert ".xlsx" in r.json()["detail"]


def test_file_rong_thi_400(with_baseline, tmp_path):
    p = tmp_path / "rong.xlsx"
    p.write_bytes(b"")
    with open(p, "rb") as fh:
        r = with_baseline.post(
            f"/api/testcases?serial={SERIAL}&package={PKG}",
            files={"file": ("rong.xlsx", fh, "application/octet-stream")},
        )
    assert r.status_code == 400
    assert "rong" in r.json()["detail"]


def test_nap_file_khi_chua_doc_baseline_thi_400(client, tmp_path):
    p = make_xlsx(tmp_path / "tc.xlsx", [
        ("1", "x", "y", "s", "", "enable_onb3_screen = false", "1. a", "1. b", "")])
    r = upload(client, p)
    assert r.status_code == 400
    assert "Chua doc baseline" in r.json()["detail"]


# ---- patch / verify / restore --------------------------------------------


def test_patch_roi_verify(with_baseline):
    ov = {"enable_onb3_screen": "false"}
    r = with_baseline.post(f"/api/patch?serial={SERIAL}&package={PKG}", json=ov)
    assert r.status_code == 200, r.text
    assert r.json()["applied"] == ov
    assert "Tat HAN app" in r.json()["note"]

    r = with_baseline.post(f"/api/verify?serial={SERIAL}&package={PKG}", json=ov)
    assert r.status_code == 200
    assert r.json()["ok"] is True
    assert r.json()["verdict_hint"] == "ok"


def test_verify_lech_thi_BLOCKED_kem_huong_dan(with_baseline):
    r = with_baseline.post(f"/api/verify?serial={SERIAL}&package={PKG}",
                           json={"enable_onb3_screen": "false"})
    body = r.json()
    assert body["ok"] is False
    assert body["verdict_hint"] == "BLOCKED"
    assert "minimumFetchInterval" in body["hint"]
    assert "KHONG phai FAIL" in body["hint"]


def test_patch_key_la_thi_400(with_baseline):
    r = with_baseline.post(f"/api/patch?serial={SERIAL}&package={PKG}",
                           json={"key_bay_dau": "true"})
    assert r.status_code == 400
    assert "khong co trong remote config" in r.json()["detail"]


def test_restore(with_baseline):
    r = with_baseline.post(f"/api/restore?serial={SERIAL}&package={PKG}")
    assert r.status_code == 200
    assert r.json()["restored"] is True

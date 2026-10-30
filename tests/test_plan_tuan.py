"""Doi chieu build voi plan tuan: workbook dung tay, khong tai sheet that."""

from __future__ import annotations

import datetime as dt

import openpyxl
import pytest

from rcr import plan_tuan as pt
from rcr import preflight

PKG = "com.aivideogenerator.aiphotoeditor.beautyeditor.imagetoimage"
NGAY = dt.date(2026, 10, 30)


def _plan(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1910-23102026"
    ws.append([14, "VSL - Inhouse", "Android", "AIP882", "- Apply SDK Next Gen"])
    ws = wb.create_sheet("2610-30102026")
    ws.append([14, "VSL - Inhouse", "Android", "AIP882", "- SDK 3.6.1", 10, "Duyên", "Testing"])
    ws.append([13, "VSL - Inhouse", "Android", "AIP882", "test v4.3.0 doing"])
    ws.append([1, "VSL - Inhouse", "Android", "AIP918", "SDK 3.7.1:\n- SDK Next Gen thêm SDK tiktok\n- Base 3.7.0"])
    ws.append([2, "VSL - Inhouse", "Android", "AIP940", "- Apply SDK Next Gen"])
    ws.append([3, "VSL - Inhouse", "Android", "AIP927", "- Dùng photo picker"])
    f = tmp_path / "plan.xlsx"
    wb.save(f)
    return f


def _kt(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "AIP882"
    wb.active.append(["Package name", PKG])
    f = tmp_path / "kt.xlsx"
    wb.save(f)
    return f


@pytest.mark.parametrize("ten, tuan", [
    ("2610-30102026", (dt.date(2026, 10, 26), dt.date(2026, 10, 30))),
    ("1910 - 23102026", (dt.date(2026, 10, 19), dt.date(2026, 10, 23))),
    ("2912-02012027", (dt.date(2026, 12, 29), dt.date(2027, 1, 2))),
    ("draf", None), ("Root", None),
])
def test_doc_tuan_tu_ten_tab(ten, tuan):
    assert pt.tuan_cua_tab(ten) == tuan


def test_lay_yeu_cau_sdk_tab_tuan_nay_bo_dong_qa(tmp_path):
    yc = pt.yeu_cau(_plan(tmp_path), "AIP882", NGAY)
    assert yc == {"tab": "2610-30102026", "sdk": ["3.6.1"], "next_gen": False, "dong": ["- SDK 3.6.1"]}
    assert pt.yeu_cau(_plan(tmp_path), "AIP882", dt.date(2026, 11, 5)) is None
    assert pt.ma_cua_package(_kt(tmp_path), PKG) == "AIP882"


def test_aip882_app_in_3_7_0_khac_plan_3_6_1(tmp_path):
    yc = pt.yeu_cau(_plan(tmp_path), "AIP882", NGAY)
    assert pt.lech(yc, "3.6.1", "3.7.0") == ["app in log SDK 3.7.0, plan ghi 3.6.1"]
    assert pt.lech(yc, "3.6.1", "3.6.1") == [] and pt.lech(yc, "", "3.6.1-alpha02") == []
    assert len(pt.lech(yc, "3.7.0", "")) == 1


def test_plan_chi_ghi_next_gen_thi_doi_ban_next_gen(tmp_path):
    yc = pt.yeu_cau(_plan(tmp_path), "AIP940", NGAY)
    assert yc["next_gen"] and not yc["sdk"]
    assert pt.lech(yc, "", "3.7.0") == [] and pt.lech(yc, "", "3.6.1") == ["app in log SDK 3.6.1, plan ghi Next-Gen"]
    yc = pt.yeu_cau(_plan(tmp_path), "AIP918", NGAY)
    assert yc["sdk"] == ["3.7.1"] and pt.lech(yc, "3.7.1-alpha01", "3.7.1") == []


def test_preflight_khac_plan_la_loi_khop_la_ok(tmp_path, monkeypatch):
    plan, kt = _plan(tmp_path), _kt(tmp_path)
    monkeypatch.setattr(preflight, "env", lambda ten: "link-plan" if ten == pt.ENV else "")
    monkeypatch.setattr(preflight.tc_source, "tai_ve", lambda link, out: kt)
    monkeypatch.setattr(pt, "tai", lambda link, out: plan)
    m = preflight.muc_plan(PKG, "3.6.1", "3.7.0", tmp_path, NGAY)
    assert m["trang_thai"] == "loi" and "KHÁC plan" in m["chi_tiet"] and "qa-auto" in m["viec_can_lam"]
    assert preflight.muc_plan(PKG, "3.6.1", "3.6.1", tmp_path, NGAY)["trang_thai"] == "ok"
    assert preflight.muc_plan(PKG, "3.6.1", "3.6.1", tmp_path, dt.date(2026, 11, 5))["trang_thai"] == "canh_bao"


def test_chua_cau_hinh_plan_chi_canh_bao(monkeypatch, tmp_path):
    monkeypatch.setattr(preflight, "env", lambda ten: "")
    assert preflight.muc_plan(PKG, "3.6.1", "3.7.0", tmp_path)["trang_thai"] == "canh_bao"
